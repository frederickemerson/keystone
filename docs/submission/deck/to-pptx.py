#!/usr/bin/env python3
"""Assemble the rendered slides into a 16:9 PowerPoint file.

Each slide is a full-bleed image: the layout, type and colour are the project's
own (docs/submission/deck/deck.html), so PowerPoint is the envelope, not the
design tool. Notes carry the spoken version of each slide.
"""
from pptx import Presentation
from pptx.util import Inches, Emu
import glob, os, sys

OUT = sys.argv[1] if len(sys.argv) > 1 else "keystone-solution-diagram.pptx"
W, H = Inches(13.333), Inches(7.5)

NOTES = {
 "s1": "Keystone. Turn retiring experts' judgement into rules AI can't override. "
       "Live at keystone-tencent.vercel.app, passcode keystone. Source is public.",
 "s2": "The judgement that keeps a building safe sits with a few people, was never written down, "
       "and leaves when they retire. A chatbot over the manual can't be trusted to act on it.",
 "s3": "The full solution diagram: problem, capture, compile, supervise, governed learning, "
       "reuse across assets, and the implementation path on Tencent Cloud.",
 "s4": "One loop, five stages, a named human at every gate. Capture, Compile, Apply, Learn, Reuse. "
       "Each stage in the guided demo is one person and one screen.",
 "s5": "How a gate decides. Trigger, then required evidence, then the action class — in that fixed "
       "order. Blocked outranks everything below it. Worked case: TS-07 at Harbourfront One, no LEW "
       "sign-off, so re-energisation is blocked and escalated to the duty LEW by name.",
 "s6": "From a spoken sentence to a released version. Read-back, draft, peer check, shadow run, "
       "approval, release — and rollback in one click. Author, shadow promoter and approver must be "
       "three different people, and the code refuses the second role and audits the refusal.",
 "s7": "An Intelligence Pill is a folder of plain text: SKILL.md, rules.yaml, evidence.yaml, "
       "provenance.json, evals and a changelog. Copied unchanged into CodeBuddy it reached the "
       "same TS-07 decision as our engine.",
 "s8": "Two paths with one line between them. Everything that decides is deterministic TypeScript; "
       "everything the model touches writes language. Hunyuan hy4-preview and MiniMax Speech 2.8 HD "
       "are integrated through Tencent TokenHub on a Singapore endpoint.",
 "s9": "Prompts elicit, extract, resolve and compose. The decision path takes no model input, "
       "so changing the model cannot change whether a gate holds.",
 "s10": "Features we would demonstrate rather than describe — teaching any pill, provenance you can "
        "hear, counterfactuals, QR capture at the plant, replay before release, asset overrides, "
        "server-side redaction, and the pill running inside CodeBuddy. And what we deliberately did not build.",
 "s11": "Five things a retrieval chatbot cannot do: a gate instead of a prompt, contradiction as "
        "arithmetic, knowledge that outlives the model, a system that refuses to invent, and "
        "governance that executes as code.",
 "s12": "What it is worth and how we would measure it: time to isolate, replay agreement, "
        "override rate, rollback frequency — measured on real incidents in a pilot.",
}

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
blank = prs.slide_layouts[6]

paths = sorted(glob.glob("slides/s*.png"), key=lambda p: int(os.path.basename(p)[1:-4]))
for p in paths:
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(p, Emu(0), Emu(0), width=W, height=H)
    key = os.path.basename(p)[:-4]
    if key in NOTES:
        slide.notes_slide.notes_text_frame.text = NOTES[key]

prs.save(OUT)
print(f"{len(paths)} slides -> {OUT} ({os.path.getsize(OUT)/1_048_576:.1f} MB)")
